from collections.abc import Generator
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.api import documents as documents_api
from app.api import folders as folders_api
from app.core.security import hash_password
from app.db.session import Base
from app.db.session import get_db as app_get_db
from app.main import app
from app.models import User
from app.services.storage import LocalStorage


@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Generator[TestClient, None, None]:
    engine = create_engine(
        f"sqlite:///{tmp_path / 'test.db'}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(bind=engine)
    testing_session = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = testing_session()
    db.add(
        User(
            username="test-admin",
            password_hash=hash_password("password-123"),
            role="admin",
            status="active",
            is_active=True,
        )
    )
    db.commit()
    db.close()

    def override_get_db():
        db = testing_session()
        try:
            yield db
        finally:
            db.close()

    testing_storage = LocalStorage(tmp_path / "storage")
    monkeypatch.setattr(documents_api, "storage", testing_storage)
    monkeypatch.setattr(folders_api, "storage", testing_storage)
    app.dependency_overrides[app_get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    engine.dispose()


def register(client: TestClient, username: str) -> tuple[str, int]:
    login = client.post(
        "/api/v1/auth/login",
        json={"username": "test-admin", "password": "password-123"},
    )
    assert login.status_code == 200
    create = client.post(
        "/api/v1/admin/users",
        headers={"Authorization": f"Bearer {login.json()['access_token']}"},
        json={"username": username, "password": "password-123"},
    )
    assert create.status_code == 201
    user_id = create.json()["id"]
    user_login = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": "password-123"},
    )
    assert user_login.status_code == 200
    return user_login.json()["access_token"], user_id


def test_public_registration_is_disabled(client: TestClient):
    response = client.post(
        "/api/v1/auth/register",
        json={"username": "用户", "password": "password-123"},
    )
    assert response.status_code in {404, 405}


def test_only_knowledge_base_owner_can_create_folder(client: TestClient):
    owner_token, _ = register(client, "folder_owner")
    viewer_token, _ = register(client, "folder_viewer")
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    viewer_headers = {"Authorization": f"Bearer {viewer_token}"}
    owner_knowledge_base_id = client.get(
        "/api/v1/knowledge-bases",
        headers=owner_headers,
    ).json()[0]["id"]

    forbidden = client.post(
        "/api/v1/folders",
        headers=viewer_headers,
        json={"knowledge_base_id": owner_knowledge_base_id, "name": "无权创建"},
    )
    assert forbidden.status_code == 403

    created = client.post(
        "/api/v1/folders",
        headers=owner_headers,
        json={"knowledge_base_id": owner_knowledge_base_id, "name": "允许创建"},
    )
    assert created.status_code == 201


def test_phase_one_document_permissions_and_lifecycle(client: TestClient):
    owner_token, owner_id = register(client, "owner_phase1")
    viewer_token, _ = register(client, "viewer_phase1")
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    viewer_headers = {"Authorization": f"Bearer {viewer_token}"}

    knowledge_bases = client.get("/api/v1/knowledge-bases", headers=owner_headers)
    assert knowledge_bases.status_code == 200
    knowledge_base_id = knowledge_bases.json()[0]["id"]

    upload = client.post(
        "/api/v1/documents/upload",
        headers=owner_headers,
        data={
            "knowledge_base_id": str(knowledge_base_id),
            "description": "phase one document",
            "tags": "制度, 测试",
        },
        files={"file": ("guide.txt", b"phase one content", "text/plain")},
    )
    assert upload.status_code == 201
    document = upload.json()
    document_id = document["id"]
    assert document["owner_id"] == owner_id
    assert document["tags"] == ["制度", "测试"]
    assert document["status"] == "ready"

    viewer_list = client.get("/api/v1/documents", headers=viewer_headers)
    assert viewer_list.status_code == 200
    assert viewer_list.json()[0]["id"] == document_id
    assert client.get(
        "/api/v1/documents?q=guide",
        headers=viewer_headers,
    ).json()[0]["id"] == document_id

    preview = client.get(f"/api/v1/documents/{document_id}/preview", headers=viewer_headers)
    assert preview.status_code == 200
    assert preview.content == b"phase one content"

    download = client.get(f"/api/v1/documents/{document_id}/download", headers=viewer_headers)
    assert download.status_code == 200
    assert download.content == b"phase one content"
    assert "attachment" in download.headers["content-disposition"]

    assert client.patch(
        f"/api/v1/documents/{document_id}",
        headers=viewer_headers,
        json={"name": "viewer-cannot-edit.txt"},
    ).status_code == 404
    assert client.delete(
        f"/api/v1/documents/{document_id}",
        headers=viewer_headers,
    ).status_code == 404

    updated = client.patch(
        f"/api/v1/documents/{document_id}",
        headers=owner_headers,
        json={"name": "renamed.txt", "tags": ["updated"]},
    )
    assert updated.status_code == 200
    assert updated.json()["name"] == "renamed.txt"
    assert updated.json()["tags"] == ["updated"]

    deleted = client.delete(f"/api/v1/documents/{document_id}", headers=owner_headers)
    assert deleted.status_code == 200
    assert deleted.json()["status"] == "deleted"
    assert client.get(f"/api/v1/documents/{document_id}", headers=viewer_headers).status_code == 404
    assert client.get(
        "/api/v1/documents?include_deleted=true",
        headers=viewer_headers,
    ).json() == []

    recycle_bin = client.get(
        "/api/v1/documents?include_deleted=true",
        headers=owner_headers,
    )
    assert recycle_bin.status_code == 200
    assert recycle_bin.json()[0]["id"] == document_id

    restored = client.post(
        f"/api/v1/documents/{document_id}/restore",
        headers=owner_headers,
    )
    assert restored.status_code == 200
    assert restored.json()["status"] == "ready"
    assert client.get(f"/api/v1/documents/{document_id}", headers=viewer_headers).status_code == 200


def test_phase_one_batch_upload(client: TestClient):
    token, _ = register(client, "batch_owner")
    headers = {"Authorization": f"Bearer {token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=headers).json()[0]["id"]

    response = client.post(
        "/api/v1/documents/upload-batch",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files=[
            ("files", ("a.md", b"# A", "text/markdown")),
            ("files", ("b.txt", b"B", "text/plain")),
        ],
    )

    assert response.status_code == 201
    assert [item["original_filename"] for item in response.json()] == ["a.md", "b.txt"]


def test_only_knowledge_base_owner_can_upload(client: TestClient):
    owner_token, _ = register(client, "upload_kb_owner")
    viewer_token, _ = register(client, "upload_kb_viewer")
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    viewer_headers = {"Authorization": f"Bearer {viewer_token}"}
    knowledge_base_id = client.get(
        "/api/v1/knowledge-bases",
        headers=owner_headers,
    ).json()[0]["id"]

    response = client.post(
        "/api/v1/documents/upload",
        headers=viewer_headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": ("blocked.txt", b"must be rejected", "text/plain")},
    )

    assert response.status_code == 403


@pytest.mark.parametrize("filename", ["slides.ppt", "slides.pptx", "report.doc", "report.docx", "table.xls", "table.xlsx"])
def test_office_document_extensions_are_supported(client: TestClient, filename: str):
    token, _ = register(client, f"office_{filename.replace('.', '_')}")
    headers = {"Authorization": f"Bearer {token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=headers).json()[0]["id"]

    response = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": (filename, b"office file placeholder", "application/octet-stream")},
    )

    assert response.status_code == 201
    assert response.json()["original_filename"] == filename
    assert response.json()["extension"] == f".{filename.rsplit('.', 1)[1]}"


def test_document_reader_handles_text_and_docx(client: TestClient):
    token, _ = register(client, "reader_owner")
    headers = {"Authorization": f"Bearer {token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=headers).json()[0]["id"]

    text_upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": ("reader.txt", "# 中文标题\n阅读内容".encode("utf-8"), "text/plain")},
    )
    assert text_upload.status_code == 201
    text_reader = client.get(f"/api/v1/documents/{text_upload.json()['id']}/reader", headers=headers)
    assert text_reader.status_code == 200
    assert text_reader.json()["kind"] == "text"
    assert text_reader.json()["pages"][0]["blocks"][0]["content"] == "# 中文标题\n阅读内容"

    word = WordDocument()
    word.add_paragraph("实验方案")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "状态"
    table.cell(1, 0).text = "解析"
    table.cell(1, 1).text = "完成"
    word_buffer = BytesIO()
    word.save(word_buffer)

    docx_upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": ("reader.docx", word_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert docx_upload.status_code == 201
    docx_reader = client.get(f"/api/v1/documents/{docx_upload.json()['id']}/reader", headers=headers)
    assert docx_reader.status_code == 200
    assert docx_reader.json()["kind"] == "docx"
    assert any(block["type"] == "table" for block in docx_reader.json()["pages"][0]["blocks"])


def test_document_reader_handles_xlsx_and_pptx(client: TestClient):
    token, _ = register(client, "reader_formats")
    headers = {"Authorization": f"Bearer {token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=headers).json()[0]["id"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["名称", "数量"])
    sheet.append(["文档", 3])
    excel_buffer = BytesIO()
    workbook.save(excel_buffer)
    xlsx_upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": ("reader.xlsx", excel_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert xlsx_upload.status_code == 201
    xlsx_reader = client.get(f"/api/v1/documents/{xlsx_upload.json()['id']}/reader", headers=headers)
    assert xlsx_reader.status_code == 200
    assert xlsx_reader.json()["kind"] == "xlsx"
    assert xlsx_reader.json()["pages"][0]["title"] == "数据"

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(914400, 914400, 4572000, 914400)
    textbox.fill.solid()
    textbox.fill.fore_color.rgb = RGBColor(242, 246, 252)
    textbox.line.color.rgb = RGBColor(31, 78, 121)
    textbox.line.width = Pt(2)
    textbox.text_frame.margin_left = Pt(12)
    textbox.text_frame.margin_right = Pt(8)
    textbox.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    textbox.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    textbox.text_frame.paragraphs[0].space_before = Pt(4)
    run = textbox.text_frame.paragraphs[0].add_run()
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    run.text = "知识库阅读器"
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, 6000000, 1000000, 1000000, 1000000)
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor(255, 230, 153)
    oval.line.color.rgb = RGBColor(192, 80, 77)
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    second_slide.shapes.add_textbox(914400, 914400, 4572000, 914400).text = "第二页"
    ppt_buffer = BytesIO()
    presentation.save(ppt_buffer)
    pptx_upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": ("reader.pptx", ppt_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert pptx_upload.status_code == 201
    pptx_reader = client.get(f"/api/v1/documents/{pptx_upload.json()['id']}/reader", headers=headers)
    assert pptx_reader.status_code == 200
    assert pptx_reader.json()["kind"] == "pptx"
    pptx_preview = client.get(f"/api/v1/documents/{pptx_upload.json()['id']}/preview", headers=headers)
    assert pptx_preview.status_code == 200
    assert pptx_preview.content.startswith(b"PK")
    assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in pptx_preview.headers["content-type"]
    page = pptx_reader.json()["pages"][0]
    assert page["shapes"]
    shape = page["shapes"][0]
    assert shape["type"] == "text"
    assert shape["paragraphs"][0]["runs"][0]["text"] == "知识库阅读器"
    assert shape["paragraphs"][0]["alignment"] == "center"
    assert shape["paragraphs"][0]["runs"][0]["style"]["font_name"] == "Microsoft YaHei"
    assert shape["paragraphs"][0]["runs"][0]["style"]["bold"] is True
    assert shape["paragraphs"][0]["runs"][0]["style"]["color"] == "#1f4e79"
    assert shape["style"]["fill"] == "#f2f6fc"
    assert shape["style"]["line"] == "#1f4e79"
    assert shape["style"]["line_width"] == 2.0
    assert shape["text_style"]["margin_left"] == 12.0
    assert shape["text_style"]["margin_right"] == 8.0
    assert shape["text_style"]["vertical_alignment"] == "middle"
    assert shape["paragraphs"][0]["style"]["space_before"] == 4.0
    assert any(item.get("geometry") == "oval" for item in page["shapes"])
    assert len(pptx_reader.json()["pages"]) == 2


def test_phase_one_folder_creation_and_document_navigation(client: TestClient, tmp_path: Path):
    token, _ = register(client, "folder_owner")
    headers = {"Authorization": f"Bearer {token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=headers).json()[0]["id"]

    create_folder = client.post(
        "/api/v1/folders",
        headers=headers,
        json={"knowledge_base_id": knowledge_base_id, "name": "实验方案"},
    )
    assert create_folder.status_code == 201
    folder = create_folder.json()
    assert folder["name"] == "实验方案"
    assert folder["parent_id"] is None
    assert folder["is_owner"] is True
    assert (tmp_path / "storage" / "folders" / str(folder["id"])).is_dir()

    create_child = client.post(
        "/api/v1/folders",
        headers=headers,
        json={
            "knowledge_base_id": knowledge_base_id,
            "parent_id": folder["id"],
            "name": "对话记录",
        },
    )
    assert create_child.status_code == 201
    child_folder_id = create_child.json()["id"]

    upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={
            "knowledge_base_id": str(knowledge_base_id),
            "folder_id": str(child_folder_id),
        },
        files={"file": ("record.md", b"# record", "text/markdown")},
    )
    assert upload.status_code == 201
    assert upload.json()["folder_id"] == child_folder_id

    child_documents = client.get(
        f"/api/v1/documents?knowledge_base_id={knowledge_base_id}&folder_id={child_folder_id}",
        headers=headers,
    )
    assert child_documents.status_code == 200
    assert len(child_documents.json()) == 1

    root_documents = client.get(
        f"/api/v1/documents?knowledge_base_id={knowledge_base_id}&root=true",
        headers=headers,
    )
    assert root_documents.status_code == 200
    assert root_documents.json() == []

    all_documents = client.get(
        f"/api/v1/documents?knowledge_base_id={knowledge_base_id}",
        headers=headers,
    )
    assert all_documents.status_code == 200
    assert [item["folder_id"] for item in all_documents.json()] == [child_folder_id]

    folders = client.get(
        f"/api/v1/folders?knowledge_base_id={knowledge_base_id}",
        headers=headers,
    )
    assert folders.status_code == 200
    assert [item["name"] for item in folders.json()] == ["实验方案", "对话记录"]

    delete_non_empty = client.delete(
        f"/api/v1/folders/{child_folder_id}",
        headers=headers,
    )
    assert delete_non_empty.status_code == 409

    empty_folder = client.post(
        "/api/v1/folders",
        headers=headers,
        json={"knowledge_base_id": knowledge_base_id, "name": "空文件夹"},
    ).json()
    assert (tmp_path / "storage" / "folders" / str(empty_folder["id"])).is_dir()
    deleted_empty = client.delete(
        f"/api/v1/folders/{empty_folder['id']}",
        headers=headers,
    )
    assert deleted_empty.status_code == 200
    assert deleted_empty.json()["status"] == "deleted"
    assert not (tmp_path / "storage" / "folders" / str(empty_folder["id"])).exists()


def test_owner_can_rename_and_move_documents_and_folders(client: TestClient):
    owner_token, _ = register(client, "move_owner")
    viewer_token, _ = register(client, "move_viewer")
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    viewer_headers = {"Authorization": f"Bearer {viewer_token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=owner_headers).json()[0]["id"]

    source = client.post(
        "/api/v1/folders",
        headers=owner_headers,
        json={"knowledge_base_id": knowledge_base_id, "name": "来源"},
    ).json()
    destination = client.post(
        "/api/v1/folders",
        headers=owner_headers,
        json={"knowledge_base_id": knowledge_base_id, "name": "目标"},
    ).json()

    upload = client.post(
        "/api/v1/documents/upload",
        headers=owner_headers,
        data={
            "knowledge_base_id": str(knowledge_base_id),
            "folder_id": str(source["id"]),
        },
        files={"file": ("move.txt", b"move me", "text/plain")},
    )
    assert upload.status_code == 201
    document_id = upload.json()["id"]

    assert client.patch(
        f"/api/v1/folders/{source['id']}",
        headers=viewer_headers,
        json={"name": "viewer-cannot-rename"},
    ).status_code == 404
    assert client.patch(
        f"/api/v1/documents/{document_id}",
        headers=viewer_headers,
        json={"folder_id": destination["id"]},
    ).status_code == 404

    renamed = client.patch(
        f"/api/v1/folders/{source['id']}",
        headers=owner_headers,
        json={"name": "来源-已重命名"},
    )
    assert renamed.status_code == 200
    assert renamed.json()["name"] == "来源-已重命名"

    moved_document = client.patch(
        f"/api/v1/documents/{document_id}",
        headers=owner_headers,
        json={"folder_id": destination["id"]},
    )
    assert moved_document.status_code == 200
    assert moved_document.json()["folder_id"] == destination["id"]

    moved_folder = client.patch(
        f"/api/v1/folders/{source['id']}",
        headers=owner_headers,
        json={"parent_id": destination["id"]},
    )
    assert moved_folder.status_code == 200
    assert moved_folder.json()["parent_id"] == destination["id"]

    cycle = client.patch(
        f"/api/v1/folders/{destination['id']}",
        headers=owner_headers,
        json={"parent_id": source["id"]},
    )
    assert cycle.status_code == 409

    moved_to_root = client.patch(
        f"/api/v1/documents/{document_id}",
        headers=owner_headers,
        json={"folder_id": None},
    )
    assert moved_to_root.status_code == 200
    assert moved_to_root.json()["folder_id"] is None


def test_markdown_preview_detects_gb18030_and_returns_utf8(client: TestClient):
    token, _ = register(client, "encoding_owner")
    headers = {"Authorization": f"Bearer {token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=headers).json()[0]["id"]

    upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": ("中文.md", "# 标题\n内容".encode("gb18030"), "text/markdown")},
    )
    assert upload.status_code == 201

    document_id = upload.json()["id"]
    preview = client.get(f"/api/v1/documents/{document_id}/preview", headers=headers)

    assert preview.status_code == 200
    assert preview.text == "# 标题\n内容"
    assert "charset=utf-8" in preview.headers["content-type"].lower()
    assert preview.headers["x-detected-encoding"] == "gb18030"


def test_html_upload_preview_and_reader_detect_encoding(client: TestClient):
    token, _ = register(client, "html_owner")
    headers = {"Authorization": f"Bearer {token}"}
    knowledge_base_id = client.get("/api/v1/knowledge-bases", headers=headers).json()[0]["id"]
    source = "<!doctype html><html><head><title>标题</title></head><body><h1>实验方案</h1><p>中文内容</p><script>bad()</script></body></html>"

    upload = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"knowledge_base_id": str(knowledge_base_id)},
        files={"file": ("实验方案.html", source.encode("gb18030"), "text/html")},
    )

    assert upload.status_code == 201
    document = upload.json()
    assert document["extension"] == ".html"
    assert document["owner_username"] == "html_owner"
    assert document["uploader_username"] == "html_owner"

    preview = client.get(f"/api/v1/documents/{document['id']}/preview", headers=headers)
    assert preview.status_code == 200
    assert "text/html" in preview.headers["content-type"]
    assert "实验方案" in preview.text
    assert preview.headers["x-detected-encoding"] == "gb18030"

    reader = client.get(f"/api/v1/documents/{document['id']}/reader", headers=headers)
    assert reader.status_code == 200
    assert reader.json()["kind"] == "html"
    content = reader.json()["pages"][0]["blocks"][0]["content"]
    assert "实验方案" in content
    assert "bad()" not in content
