import base64
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from app.services.text_encoding import decode_text


class DocumentReaderError(Exception):
    pass


TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
HTML_EXTENSIONS = {".html", ".htm"}
IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}
THEME_COLOR_DEFAULTS = {
    "DARK_1": "#000000",
    "LIGHT_1": "#ffffff",
    "DARK_2": "#44546a",
    "LIGHT_2": "#e7e6e6",
    "ACCENT_1": "#4472c4",
    "ACCENT_2": "#ed7d31",
    "ACCENT_3": "#a5a5a5",
    "ACCENT_4": "#ffc000",
    "ACCENT_5": "#5b9bd5",
    "ACCENT_6": "#70ad47",
    "HYPERLINK": "#0563c1",
    "FOLLOWED_HYPERLINK": "#954f72",
}


def _text_page(content: str) -> dict:
    return {"number": 1, "blocks": [{"type": "text", "content": content}]}


def _read_text(path: Path) -> dict:
    decoded = decode_text(path.read_bytes())
    return {
        "kind": "text",
        "encoding": decoded.encoding,
        "pages": [_text_page(decoded.content)],
    }


class _HtmlTextExtractor(HTMLParser):
    _ignored_tags = {"head", "script", "style", "noscript", "template", "svg"}
    _block_tags = {
        "address",
        "article",
        "aside",
        "blockquote",
        "br",
        "div",
        "footer",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "li",
        "main",
        "nav",
        "ol",
        "p",
        "pre",
        "section",
        "table",
        "td",
        "th",
        "tr",
        "ul",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        normalized = tag.lower()
        if normalized in self._ignored_tags:
            self._ignored_depth += 1
            return
        if self._ignored_depth == 0 and normalized in self._block_tags:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.lower()
        if normalized in self._ignored_tags and self._ignored_depth:
            self._ignored_depth -= 1
            return
        if self._ignored_depth == 0 and normalized in self._block_tags:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._ignored_depth == 0:
            self.parts.append(data)


def _read_html(path: Path) -> dict:
    decoded = decode_text(path.read_bytes())
    parser = _HtmlTextExtractor()
    parser.feed(decoded.content)
    parser.close()
    text = re.sub(r"[ \t]+\n", "\n", "".join(parser.parts))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return {
        "kind": "html",
        "encoding": decoded.encoding,
        "pages": [_text_page(text)],
    }


def _read_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(
            {
                "number": index,
                "blocks": [{"type": "text", "content": page.extract_text() or ""}],
            }
        )
    return {"kind": "pdf", "encoding": "utf-8", "pages": pages}


def _read_docx(path: Path) -> dict:
    document = WordDocument(str(path))
    blocks = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append({"type": "paragraph", "content": paragraph.text})
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        if rows:
            blocks.append({"type": "table", "rows": rows})
    return {"kind": "docx", "encoding": "utf-8", "pages": [{"number": 1, "blocks": blocks}]}


def _read_xlsx(path: Path) -> dict:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    rows.append(values)
            pages.append(
                {
                    "number": index,
                    "title": sheet.title,
                    "blocks": [{"type": "table", "rows": rows}] if rows else [],
                }
            )
    finally:
        workbook.close()
    return {"kind": "xlsx", "encoding": "utf-8", "pages": pages}


def _pptx_color(color) -> str | None:
    try:
        rgb = color.rgb
        return f"#{rgb}".lower()
    except (AttributeError, TypeError, ValueError):
        try:
            theme_color = color.theme_color
            name = getattr(theme_color, "name", str(theme_color).split(".")[-1])
            return THEME_COLOR_DEFAULTS.get(name)
        except (AttributeError, TypeError, ValueError):
            return None


def _pptx_fill(shape) -> str | None:
    try:
        return _pptx_color(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def _pptx_line(shape) -> str | None:
    try:
        return _pptx_color(shape.line.color)
    except (AttributeError, TypeError, ValueError):
        return None


def _pptx_run_style(run) -> dict:
    font = run.font
    style = {
        "font_name": font.name,
        "font_size": font.size.pt if font.size else None,
        "bold": bool(font.bold),
        "italic": bool(font.italic),
        "underline": bool(font.underline),
        "color": _pptx_color(font.color),
    }
    return {key: value for key, value in style.items() if value is not None}


def _pptx_alignment(alignment) -> str | None:
    if alignment is None:
        return None
    value = str(alignment).lower()
    for name in ("center", "right", "justify", "left"):
        if name in value:
            return name
    return None


def _pptx_vertical_alignment(anchor) -> str | None:
    if anchor is None:
        return None
    value = str(anchor).lower()
    for name in ("top", "middle", "bottom"):
        if name in value:
            return name
    return None


def _pptx_text_paragraphs(shape) -> list[dict]:
    paragraphs = []
    for paragraph in shape.text_frame.paragraphs:
        runs = [
            {"text": run.text, "style": _pptx_run_style(run)}
            for run in paragraph.runs
            if run.text
        ]
        if not runs and paragraph.text:
            runs = [{"text": paragraph.text, "style": {}}]
        if runs:
            paragraph_style = {}
            if paragraph.space_before is not None:
                paragraph_style["space_before"] = paragraph.space_before.pt
            if paragraph.space_after is not None:
                paragraph_style["space_after"] = paragraph.space_after.pt
            if paragraph.line_spacing is not None:
                if hasattr(paragraph.line_spacing, "pt"):
                    paragraph_style["line_spacing_pt"] = paragraph.line_spacing.pt
                else:
                    paragraph_style["line_spacing_multiple"] = float(paragraph.line_spacing)
            paragraphs.append(
                {
                    "level": paragraph.level,
                    "alignment": _pptx_alignment(paragraph.alignment),
                    "runs": runs,
                    "style": paragraph_style,
                }
            )
    return paragraphs


def _pptx_text_style(shape) -> dict:
    text_frame = shape.text_frame
    style = {
        "margin_left": text_frame.margin_left.pt,
        "margin_right": text_frame.margin_right.pt,
        "margin_top": text_frame.margin_top.pt,
        "margin_bottom": text_frame.margin_bottom.pt,
        "vertical_alignment": _pptx_vertical_alignment(text_frame.vertical_anchor),
        "word_wrap": text_frame.word_wrap,
    }
    return {key: value for key, value in style.items() if value is not None}


def _pptx_shape_style(shape) -> dict:
    style = {
        "fill": _pptx_fill(shape),
        "line": _pptx_line(shape),
        "line_width": shape.line.width.pt if shape.line.width else None,
        "rotation": float(getattr(shape, "rotation", 0) or 0),
    }
    return {key: value for key, value in style.items() if value is not None}


def _pptx_shape_geometry(shape) -> str | None:
    try:
        value = str(shape.auto_shape_type).lower()
    except (AttributeError, TypeError, ValueError):
        return None
    if "oval" in value:
        return "oval"
    if "round" in value:
        return "rounded"
    if "line" in value:
        return "line"
    return None


def _pptx_background_color(slide) -> str | None:
    layout = getattr(slide, "slide_layout", None)
    master = getattr(layout, "slide_master", None) if layout is not None else None
    candidates = [
        slide.background,
        layout.background if layout is not None else None,
        master.background if master is not None else None,
    ]
    for background in candidates:
        if background is None:
            continue
        try:
            color = _pptx_color(background.fill.fore_color)
        except (AttributeError, TypeError, ValueError):
            color = None
        if color:
            return color
        try:
            bg = background._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
            scheme = None
            srgb = None
            if bg is not None:
                for child in bg.iter():
                    if child.tag == "{http://schemas.openxmlformats.org/drawingml/2006/main}schemeClr":
                        scheme = child
                    if child.tag == "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr":
                        srgb = child
            if scheme is not None:
                scheme_name = {
                    "bg1": "LIGHT_1",
                    "bg2": "LIGHT_2",
                    "tx1": "DARK_1",
                    "tx2": "DARK_2",
                }.get(scheme.get("val"))
                if scheme_name:
                    return THEME_COLOR_DEFAULTS[scheme_name]
            if srgb is not None and srgb.get("val"):
                return f"#{srgb.get('val')}".lower()
        except (AttributeError, TypeError, ValueError):
            pass
    return None


def _pptx_shape_record(shape, slide_width: int, slide_height: int) -> dict | None:
    record = {
        "left": max(0, min(1, float(shape.left) / slide_width)),
        "top": max(0, min(1, float(shape.top) / slide_height)),
        "width": max(0, min(1, float(shape.width) / slide_width)),
        "height": max(0, min(1, float(shape.height) / slide_height)),
        "style": _pptx_shape_style(shape),
    }
    geometry = _pptx_shape_geometry(shape)
    if geometry:
        record["geometry"] = geometry
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
        image = shape.image
        encoded = base64.b64encode(image.blob).decode("ascii")
        return {
            **record,
            "type": "image",
            "data_uri": f"data:{image.content_type};base64,{encoded}",
        }
    if getattr(shape, "has_table", False):
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        return {**record, "type": "table", "rows": rows}
    if getattr(shape, "has_text_frame", False) and shape.text.strip():
        paragraphs = _pptx_text_paragraphs(shape)
        return {
            **record,
            "type": "text",
            "paragraphs": paragraphs,
            "text_style": _pptx_text_style(shape),
        }
    if record["style"].get("fill") or record["style"].get("line"):
        return {**record, "type": "shape"}
    return None


def _pptx_shape_records(shapes, slide_width: int, slide_height: int) -> list[dict]:
    records = []
    for shape in shapes:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            records.extend(_pptx_shape_records(shape.shapes, slide_width, slide_height))
            continue
        record = _pptx_shape_record(shape, slide_width, slide_height)
        if record:
            records.append(record)
    return records


def _read_pptx(path: Path) -> dict:
    presentation = Presentation(str(path))
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    pages = []
    for index, slide in enumerate(presentation.slides, start=1):
        page = {
            "number": index,
            "width": slide_width,
            "height": slide_height,
            "shapes": _pptx_shape_records(slide.shapes, slide_width, slide_height),
        }
        background = _pptx_background_color(slide)
        if background:
            page["background"] = background
        pages.append(page)
    return {"kind": "pptx", "encoding": "utf-8", "pages": pages}


def _read_legacy_office(path: Path) -> dict:
    try:
        import olefile

        if not olefile.isOleFile(str(path)):
            raise DocumentReaderError("无法解析该旧版 Office 文件")
        ole = olefile.OleFileIO(str(path))
        chunks = []
        try:
            for stream_path in ole.listdir():
                data = ole.openstream(stream_path).read()
                for encoding in ("utf-16-le", "utf-8"):
                    text = data.decode(encoding, errors="ignore")
                    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]+", " ", text)
                    text = re.sub(r"[^\w\s，。！？、；：（）《》“”‘’【】,.!?;:()[\]{}<>%+\-*/=@#_]", "", text)
                    if len(text.strip()) >= 8:
                        chunks.append(text.strip())
                        break
        finally:
            ole.close()
        content = "\n\n".join(dict.fromkeys(chunks))
        if not content:
            raise DocumentReaderError("旧版 Office 文件没有可读取的文本内容")
        return {"kind": "legacy-office", "encoding": "utf-8", "pages": [_text_page(content)]}
    except DocumentReaderError:
        raise
    except Exception as exc:
        raise DocumentReaderError(f"旧版 Office 文件解析失败：{exc}") from exc


def _read_image(path: Path, extension: str) -> dict:
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return {
        "kind": "image",
        "encoding": "binary",
        "data_uri": f"data:{IMAGE_MIME_TYPES[extension]};base64,{encoded}",
        "pages": [],
    }


def read_document(path: Path, extension: str) -> dict:
    if not path.is_file():
        raise DocumentReaderError("Stored file not found")
    try:
        if extension in TEXT_EXTENSIONS:
            return _read_text(path)
        if extension in HTML_EXTENSIONS:
            return _read_html(path)
        if extension == ".pdf":
            return _read_pdf(path)
        if extension == ".docx":
            return _read_docx(path)
        if extension == ".xlsx":
            return _read_xlsx(path)
        if extension == ".pptx":
            return _read_pptx(path)
        if extension in IMAGE_MIME_TYPES:
            return _read_image(path, extension)
        if extension in {".doc", ".xls", ".ppt"}:
            return _read_legacy_office(path)
    except DocumentReaderError:
        raise
    except Exception as exc:
        raise DocumentReaderError(f"文档解析失败：{exc}") from exc
    raise DocumentReaderError(f"暂不支持在阅读器中打开 {extension or '该类型'} 文件")
